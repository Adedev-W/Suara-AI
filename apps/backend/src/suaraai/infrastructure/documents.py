from __future__ import annotations

import re
from io import BytesIO

from pptx import Presentation
from pypdf import PdfReader

from suaraai.domain.copilot import ExtractedDocument


class DocumentParseError(ValueError):
    pass


def parse_document(source_name: str, content_type: str | None, data: bytes) -> ExtractedDocument:
    lowered_name = source_name.lower()
    if lowered_name.endswith(".pdf") or content_type == "application/pdf":
        return _parse_pdf(source_name, data)
    if lowered_name.endswith((".pptx", ".ppt")) or content_type in {
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "application/vnd.ms-powerpoint",
    }:
        if lowered_name.endswith(".ppt") or content_type == "application/vnd.ms-powerpoint":
            raise DocumentParseError("Legacy .ppt files are not supported; upload .pptx")
        return _parse_pptx(source_name, data)
    raise DocumentParseError("Only PDF and PPTX documents are supported")


def chunk_document(
    document: ExtractedDocument, max_chars: int = 1200
) -> list[tuple[int | None, str]]:
    chunks: list[tuple[int | None, str]] = []
    for page_number, page_text in document.pages:
        words = re.findall(r"\S+", page_text)
        current: list[str] = []
        current_length = 0
        for word in words:
            if current and current_length + len(word) + 1 > max_chars:
                chunks.append((page_number, " ".join(current)))
                overlap = current[-25:]
                current = overlap
                current_length = len(" ".join(current))
            current.append(word)
            current_length += len(word) + 1
        if current:
            chunks.append((page_number, " ".join(current)))
    return chunks


def _parse_pdf(source_name: str, data: bytes) -> ExtractedDocument:
    try:
        reader = PdfReader(BytesIO(data))
        pages: list[tuple[int | None, str]] = []
        for index, page in enumerate(reader.pages):
            extracted_text = page.extract_text()
            pages.append((index + 1, extracted_text if isinstance(extracted_text, str) else ""))
    except Exception as exc:
        raise DocumentParseError("Could not read PDF document") from exc
    return _with_nonempty_pages(source_name, pages)


def _parse_pptx(source_name: str, data: bytes) -> ExtractedDocument:
    try:
        presentation = Presentation(BytesIO(data))
        pages: list[tuple[int | None, str]] = []
        for index, slide in enumerate(presentation.slides):
            text = " ".join(
                shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text
            )
            pages.append((index + 1, text))
    except Exception as exc:
        raise DocumentParseError("Could not read PPTX document") from exc
    return _with_nonempty_pages(source_name, pages)


def _with_nonempty_pages(
    source_name: str, pages: list[tuple[int | None, str]]
) -> ExtractedDocument:
    normalized = [(page, text.strip()) for page, text in pages if text.strip()]
    if not normalized:
        raise DocumentParseError("Document contains no readable text")
    return ExtractedDocument(source_name=source_name, pages=normalized)


class LocalDocumentParser:
    def parse(self, source_name: str, content_type: str | None, data: bytes) -> ExtractedDocument:
        return parse_document(source_name, content_type, data)

    def chunk(self, document: ExtractedDocument) -> list[tuple[int | None, str]]:
        return chunk_document(document)
